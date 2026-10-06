"""Derive the S5 inputs and validation targets in data/derived/ from the raw RxCADRE downloads.

Outputs:
- s5_perimeter_xy.json        burn-block polygon, UTM 16N (NAD83)
- s5_fire_sensors_located.csv fire-behavior package arrival times + surveyed positions
- s5_anemometers.csv          1 Hz anemometer records inside S5
- s5_measurements.json        measured rate of spread, wind during the burn, ignition estimate
"""

from __future__ import annotations

import json
import re

import numpy as np
import openpyxl
import pandas as pd
import shapefile

from rxcadre_sim.paths import DERIVED_DIR, RAW_DIR
from rxcadre_sim.s5_case import ignition_centre_utm

MPH_TO_M_S = 0.44704
WIND_WINDOW_BEFORE_S = 1800.0
WIND_STEP_S = 60.0
WIND_SERIES_DURATION_S = 2400.0
# The consistent S5 anemometers agree within ±20°; the misaligned ones sit 40–71° away.
MAX_DIRECTION_OFFSET_DEG = 30.0


def s5_perimeter() -> list[tuple[float, float]]:
    reader = shapefile.Reader(
        str(RAW_DIR / "burn_blocks/2012_fire_perimeters/2012_burn_block_perimeters.shp")
    )
    for record, shape in zip(reader.records(), reader.shapes(), strict=True):
        if record.as_dict()["Block_Name"].strip() == "S5":
            return [tuple(point) for point in shape.points]
    raise LookupError("S5 not found in burn-block perimeters")


def fire_sensor_summaries() -> pd.DataFrame:
    """One row per S5 fire-behavior package from each workbook's summary sheet."""
    rows = []
    for workbook_path in sorted((RAW_DIR / "fire_behavior/FireBehavior/s5").glob("*.xlsx")):
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        summary = {}
        for row in workbook["summary"].iter_rows(values_only=True):
            values = [v for v in row if v is not None]
            if len(values) >= 2:
                summary[str(values[0]).strip()] = values[1]
        rows.append(
            {
                "tag": int(re.search(r"t(\d+)", workbook_path.name).group(1)),
                "logger": summary.get("Logger Number"),
                "arrival_utc": str(summary.get("Time of NAR Signal Detection")),
                "peak_total_kw_m2": summary.get("Peak Total Flux(kW/m^2)"),
                "peak_radiant_kw_m2": summary.get("Peak Radiant Flux(kW/m^2)"),
                "integrated_total_kj_m2": summary.get("Integrated Total Flux(kW*s/m^2)"),
            }
        )
    return pd.DataFrame(rows)


def survey_points() -> pd.DataFrame:
    reader = shapefile.Reader(
        str(RAW_DIR / "survey/SurveyPointInformation.shp"), encoding="latin-1"
    )
    points = pd.DataFrame([record.as_dict() for record in reader.records()])
    points["x"] = [shape.points[0][0] for shape in reader.shapes()]
    points["y"] = [shape.points[0][1] for shape in reader.shapes()]
    points["tag"] = pd.to_numeric(points.Tag_Number, errors="coerce")
    return points[points.Burn_Block.astype(str).str.strip() == "S5"]


def located_fire_sensors() -> pd.DataFrame:
    """Merge arrival times with surveyed positions and flag the package whose clock failed."""
    sensors = fire_sensor_summaries().merge(
        survey_points()[["tag", "Sensor_ID", "x", "y", "Ellipsoid_"]], on="tag", how="left"
    )
    arrival_s = pd.to_timedelta(sensors.arrival_utc).dt.total_seconds()
    median_s = arrival_s.median()
    sensors["is_valid"] = (arrival_s - median_s).abs() < 3600
    return sensors.sort_values("arrival_utc")


def s5_anemometers() -> pd.DataFrame:
    columns = {
        "_Time_GMT_00_00": "time",
        "_WindSpeed_mph": "ws_mph",
        "_WindGust_mph": "gust_mph",
        "_WindDirection_Degrees_from_nort": "wd",
        "_PlotID": "plot_id",
        "_InstrumentID": "inst",
        "_Easting__X_": "E",
        "_Northing__Y_": "N",
    }
    records = pd.read_csv(
        RAW_DIR / "fire_behavior/Anemometer/Anemometer_Final_v2_only_Splots.csv",
        encoding="latin-1",
        usecols=list(columns),
        low_memory=False,
    ).rename(columns=columns)
    records = records[records.plot_id.astype(str).str.strip().str.upper() == "S5"]
    return records[(records.ws_mph >= 0) & (records.wd >= 0)]


def measured_spread(sensors: pd.DataFrame) -> dict[str, float]:
    """Plane fit of arrival time over position.

    ROS = 1/|grad t|; the intercept is the extrapolated ignition time at the line centre.
    """
    valid = sensors[sensors.is_valid]
    arrival_s = pd.to_timedelta(valid.arrival_utc).dt.total_seconds().to_numpy()
    design = np.c_[
        np.ones(len(valid)), valid.x - ignition_centre_utm()[0], valid.y - ignition_centre_utm()[1]
    ]
    coefficients, *_ = np.linalg.lstsq(design, arrival_s, rcond=None)
    return {
        "ros_m_s": float(1.0 / np.hypot(*coefficients[1:])),
        "spread_heading_deg": float(np.degrees(np.arctan2(coefficients[1], coefficients[2])) % 360),
        "ignition_estimate_utc_s": float(coefficients[0]),
        "first_arrival_utc_s": float(arrival_s.min()),
        "last_arrival_utc_s": float(arrival_s.max()),
    }


def _seconds(anemometers: pd.DataFrame) -> pd.Series:
    return pd.to_timedelta(anemometers.time.str.strip()).dt.total_seconds()


def _vector_mean_from_deg(window: pd.DataFrame) -> float:
    """Direction of the average wind arrow (a plain average fails across north: 350° and 10°)."""
    speed = window.ws_mph * MPH_TO_M_S
    from_rad = np.radians(window.wd)
    east, north = -(speed * np.sin(from_rad)).mean(), -(speed * np.cos(from_rad)).mean()
    return float(np.degrees(np.arctan2(-east, -north)) % 360)


def misaligned_anemometers(anemometers: pd.DataFrame, start_s: float, end_s: float) -> list[str]:
    """Instruments whose mean direction is more than MAX_DIRECTION_OFFSET_DEG from the group.

    The RxCADRE metadata warns that some anemometers were misaligned: their direction is wrong
    but their speed is correct. Their speeds are still used; their directions are not.
    """
    seconds = _seconds(anemometers)
    window = anemometers[(seconds >= start_s) & (seconds <= end_s)]
    directions = window.groupby("inst").apply(_vector_mean_from_deg, include_groups=False)
    # Repeat until stable, so outliers do not drag the group direction toward each other.
    excluded: set[str] = set()
    while True:
        kept = directions.drop(sorted(excluded))
        unit = np.c_[np.sin(np.radians(kept)), np.cos(np.radians(kept))]
        group = np.degrees(np.arctan2(*unit.mean(axis=0))) % 360
        offset = (directions - group + 180) % 360 - 180
        flagged = set(directions.index[offset.abs() > MAX_DIRECTION_OFFSET_DEG])
        if flagged == excluded:
            return sorted(excluded)
        excluded = flagged


def wind_during(
    anemometers: pd.DataFrame, start_s: float, end_s: float, *, exclude_direction: list[str]
) -> dict[str, float]:
    seconds = _seconds(anemometers)
    window = anemometers[(seconds >= start_s) & (seconds <= end_s)]
    aligned = window[~window.inst.isin(exclude_direction)]
    speed = window.ws_mph * MPH_TO_M_S
    aligned_speed = aligned.ws_mph * MPH_TO_M_S
    from_rad = np.radians(aligned.wd)
    east = -(aligned_speed * np.sin(from_rad)).mean()
    north = -(aligned_speed * np.cos(from_rad)).mean()
    return {
        "mean_speed_m_s": float(speed.mean()),
        "gust_p95_m_s": float((window.gust_mph * MPH_TO_M_S).quantile(0.95)),
        "vector_mean_from_deg": _vector_mean_from_deg(aligned),
        "steadiness": float(np.hypot(east, north) / aligned_speed.mean()),
        "samples": int(len(window)),
        "direction_from_instruments": sorted(set(aligned.inst)),
    }


def wind_series(
    anemometers: pd.DataFrame, start_s: float, duration_s: float, *, exclude_direction: list[str]
) -> dict[str, list[float]]:
    """Minute-by-minute wind from start_s: speed from all anemometers, direction from aligned."""
    seconds = _seconds(anemometers)
    window = anemometers[(seconds >= start_s) & (seconds < start_s + duration_s)].copy()
    window["step"] = ((seconds[window.index] - start_s) // WIND_STEP_S).astype(int)
    speed = window.groupby("step").ws_mph.mean() * MPH_TO_M_S
    aligned = window[~window.inst.isin(exclude_direction)]
    direction = aligned.groupby("step").apply(_vector_mean_from_deg, include_groups=False)
    steps = speed.index.intersection(direction.index)
    return {
        "times_s": [float(step * WIND_STEP_S) for step in steps],
        "speed_m_s": [round(float(speed[step]), 3) for step in steps],
        "from_deg": [round(float(direction[step]), 1) for step in steps],
    }


def main() -> None:
    DERIVED_DIR.mkdir(parents=True, exist_ok=True)
    (DERIVED_DIR / "s5_perimeter_xy.json").write_text(json.dumps(s5_perimeter()))

    sensors = located_fire_sensors()
    sensors.to_csv(DERIVED_DIR / "s5_fire_sensors_located.csv", index=False)

    anemometers = s5_anemometers()
    anemometers.to_csv(DERIVED_DIR / "s5_anemometers.csv", index=False)

    spread = measured_spread(sensors)
    start, end = spread["ignition_estimate_utc_s"], spread["last_arrival_utc_s"]
    misaligned = misaligned_anemometers(anemometers, start, end)
    measurements = {
        "spread": spread,
        "misaligned_anemometers": misaligned,
        "wind_during_burn": wind_during(anemometers, start, end, exclude_direction=misaligned),
        "wind_before_burn": wind_during(
            anemometers, start - WIND_WINDOW_BEFORE_S, start, exclude_direction=misaligned
        ),
        "wind_series_after_ignition": wind_series(
            anemometers, start, WIND_SERIES_DURATION_S, exclude_direction=misaligned
        ),
    }
    (DERIVED_DIR / "s5_measurements.json").write_text(json.dumps(measurements, indent=2))
    print(json.dumps(measurements, indent=2))
    print(f"wrote derived S5 data to {DERIVED_DIR}")


if __name__ == "__main__":
    main()
